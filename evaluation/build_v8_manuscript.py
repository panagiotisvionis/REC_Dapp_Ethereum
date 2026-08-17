"""
build_v8_manuscript.py — v2 (corrected)
========================================
Builds RecChain_manuscript_v8_complete.docx from v7 by:
  1. Inserting Figure 1 (replacing INSERT placeholder)
  2. Inserting Figure 2 (replacing INSERT placeholder)
  3. Replacing Figure 6 image (old wrong sensitivity chart → corrected one)
  4. Correcting Table IV with real sensitivity-analysis values
  5. Inserting GitHub URL
  6. Fixing ROC-AUC text reference (1.0000 → 0.9999)
  7. Removing "→ BEFORE SUBMISSION" editor note

Note: Figures 3,4,5,7 are already correctly embedded in v7 — left untouched.

Output: /home/lysos/Desktop/Panagiotis/papers/RecChain/RecChain_manuscript_v8_complete.docx
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import copy

# ── Paths ─────────────────────────────────────────────────────────────────────
SRC   = Path("/home/lysos/Desktop/Panagiotis/papers/RecChain/RecChain_manuscript_v7_FIGURES.docx")
OUT   = Path("/home/lysos/Desktop/Panagiotis/papers/RecChain/RecChain_manuscript_v8_complete.docx")
FIGS  = Path(__file__).parent / "results" / "figures"
GITHUB_URL = "https://github.com/panagiotisvionis/REC_Dapp_Ethereum"

CAPTION_FIG1 = "Figure 1. RecChain four-layer architecture. Dashed arrow = direct MetaMask wallet interaction (no API required)."
CAPTION_FIG2 = "Figure 2. Oracle-based certificate issuance sequence. Solid arrows = calls/requests; dashed arrows = responses."

# Correct Table IV values (seed=42, n=1000, 10% anomaly, ENTSO-E params)
TABLE_IV_DATA = [
    # threshold, precision, recall, f1, fpr, roc_auc, fn
    ("≥ 2.0",   "0.7246", "1.0000", "0.8403", "0.0422", "0.9999", "0"),
    ("≥ 2.5 ★", "0.9009", "1.0000", "0.9479", "0.0122", "0.9999", "0"),
    ("≥ 3.0",   "0.9897", "0.9600", "0.9746", "0.0011", "0.9999", "4"),
    ("≥ 3.5",   "1.0000", "0.9200", "0.9583", "0.0000", "0.9999", "8"),
    ("≥ 4.0",   "1.0000", "0.7800", "0.8764", "0.0000", "0.9999", "22"),
]

# ── Helpers ───────────────────────────────────────────────────────────────────

def center_para(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def make_image_para(doc, png_path, width_in):
    p = doc.add_paragraph()
    center_para(p)
    p.add_run().add_picture(str(png_path), width=Inches(width_in))
    return p

def make_caption_para(doc, caption_text):
    p = doc.add_paragraph()
    center_para(p)
    run = p.add_run(caption_text)
    run.italic = True
    run.font.size = Pt(10)
    return p

def insert_before(anchor_el, *new_els):
    for el in new_els:
        anchor_el.addprevious(el)

def replace_placeholder_with_image(placeholder_para, doc, png_path, width_in, caption):
    """
    Turn a placeholder text paragraph into a centred image paragraph,
    then append a caption paragraph immediately after it.
    """
    # Repurpose the paragraph: clear all runs, centre it, add picture
    for child in list(placeholder_para._element):
        if child.tag.endswith('}r') or child.tag.endswith('}hyperlink'):
            placeholder_para._element.remove(child)
    center_para(placeholder_para)
    placeholder_para.add_run().add_picture(str(png_path), width=Inches(width_in))

    # Insert caption after (addnext inserts at the BEGINNING of "next" siblings,
    # so we use addnext on the para element itself)
    cap_para = make_caption_para(doc, caption)
    placeholder_para._element.addnext(cap_para._element)

def replace_existing_image_para(image_para, doc, png_path, width_in):
    """
    Replace the picture in an existing image paragraph with a new one.
    image_para contains a run with an inline picture.
    """
    # Find the drawing element and remove it
    for child in list(image_para._element.iter()):
        if child.tag.endswith('}drawing'):
            child.getparent().remove(child)
            break
    # Repurpose the first run or add a new one
    existing_runs = image_para._element.findall('.//' + qn('w:r'))
    if existing_runs:
        # Remove all existing runs
        for r in existing_runs:
            r.getparent().remove(r)
    center_para(image_para)
    image_para.add_run().add_picture(str(png_path), width=Inches(width_in))

def set_cell_text(cell, text, bold=False):
    para = cell.paragraphs[0]
    para.clear()
    run = para.add_run(text)
    run.bold = bold

# ── Load document ─────────────────────────────────────────────────────────────
doc = Document(str(SRC))
paras = doc.paragraphs

# ── Find target paragraphs ────────────────────────────────────────────────────
insert_fig1  = next(p for p in paras if "INSERT Figure 1" in p.text)
insert_fig2  = next(p for p in paras if "INSERT Figure 2" in p.text)
before_sub   = next((p for p in paras if "BEFORE SUBMISSION" in p.text), None)
para_github  = next(p for p in paras if "[INSERT GitHub URL]" in p.text)
para_roc_ref = next((p for p in paras if "ROC-AUC = 1.0000" in p.text), None)

# Find Figure 6's existing image paragraph (the one BEFORE the Fig 6 caption)
caption_fig6 = next(p for p in paras if p.text.startswith("Figure 6."))
# Walk backwards to find the image paragraph
fig6_img_para = None
idx6 = paras.index(caption_fig6)
for i in range(idx6 - 1, max(0, idx6 - 5), -1):
    candidate = paras[i]
    blips = candidate._element.findall(
        './/' + qn('a:blip'),
        namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    )
    if blips:
        fig6_img_para = candidate
        break

print(f"INSERT Fig1 para: '{insert_fig1.text[:60]}'")
print(f"INSERT Fig2 para: '{insert_fig2.text[:60]}'")
print(f"Fig6 image para index: {paras.index(fig6_img_para) if fig6_img_para else 'NOT FOUND'}")

# ── 1. Insert Figure 1 (replace placeholder) ──────────────────────────────────
replace_placeholder_with_image(
    insert_fig1, doc,
    FIGS / "fig1_architecture.png", 5.5, CAPTION_FIG1
)
print("Figure 1 inserted ✓")

# ── 2. Insert Figure 2 (replace placeholder) ──────────────────────────────────
replace_placeholder_with_image(
    insert_fig2, doc,
    FIGS / "fig2_oracle_sequence.png", 5.5, CAPTION_FIG2
)
print("Figure 2 inserted ✓")

# ── 3. Replace Figure 6 with corrected sensitivity chart ─────────────────────
if fig6_img_para:
    replace_existing_image_para(fig6_img_para, doc, FIGS / "fig6_sensitivity_chart.png", 5.5)
    print("Figure 6 replaced with corrected sensitivity chart ✓")
else:
    print("WARNING: Figure 6 image paragraph not found — inserting before caption")
    img_p = make_image_para(doc, FIGS / "fig6_sensitivity_chart.png", 5.5)
    caption_fig6._element.addprevious(img_p._element)

# ── 4. Correct Table IV ───────────────────────────────────────────────────────
tbl4 = doc.tables[2]
for row_idx, row_data in enumerate(TABLE_IV_DATA, start=1):
    thresh, prec, rec, f1, fpr, auc, fn = row_data
    cells = tbl4.rows[row_idx].cells
    set_cell_text(cells[0], thresh, bold=(row_idx == 2))
    set_cell_text(cells[1], prec)
    set_cell_text(cells[2], rec)
    set_cell_text(cells[3], f1)
    set_cell_text(cells[4], fpr)
    set_cell_text(cells[5], auc)
    set_cell_text(cells[6], fn)
print("Table IV corrected ✓")

# ── 5. Insert GitHub URL ──────────────────────────────────────────────────────
full_text = para_github.text
if "[INSERT GitHub URL]" in full_text:
    para_github.clear()
    para_github.add_run(full_text.replace("[INSERT GitHub URL]", GITHUB_URL))
print(f"GitHub URL inserted ✓")

# ── 6. Fix ROC-AUC text reference ────────────────────────────────────────────
if para_roc_ref:
    old = para_roc_ref.text
    new = old.replace("ROC-AUC = 1.0000", "ROC-AUC = 0.9999")
    para_roc_ref.clear()
    para_roc_ref.add_run(new)
    print("ROC-AUC text fixed ✓")

# ── 7. Remove editor note ─────────────────────────────────────────────────────
if before_sub:
    before_sub._element.getparent().remove(before_sub._element)
    print("Editor note removed ✓")

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save(str(OUT))
import os
size = os.path.getsize(OUT) // 1024
print(f"\n✓ Saved: {OUT.name}  ({size} KB)")
print(f"  Inline images: {len(doc.inline_shapes)}")
print(f"  Tables: {len(doc.tables)}")
