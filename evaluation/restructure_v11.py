"""
restructure_v11.py
==================
Takes v10 (manual edits preserved) and restructures from 11 → 7 chapters.

Merges:
  4. Smart Contract Design
  5. Oracle-Based Certificate Issuance
  → 4. Smart Contract Design and Oracle Integration

  6. Intelligent Anomaly Detection Module
  7. Evaluation
  → 5. Anomaly Detection Module and Evaluation
    5.1 Detection Methodology and Decision Logic   (was 6.1)
    5.2 Anomaly Detection Evaluation              (was 7.1)
    5.3 Sensitivity Analysis                      (was 7.2)
    5.4 Gas Cost Analysis                         (was 7.3)

  8. Implementation
  9. Security Analysis
  10. Limitations and Future Work
  → 6. Implementation, Security Analysis, and Limitations
    6.1 Evaluation Scope and Detection Methodology Limitations  (was 10.1)
    6.2 Deployment Constraints and Future Research Directions   (was 10.2)

  11. Conclusion → 7. Conclusion

Output: RecChain_manuscript_v11.docx
"""

from pathlib import Path
from docx import Document
from docx.shared import RGBColor

SRC = Path("/home/lysos/Desktop/Panagiotis/papers/RecChain/RecChain_manuscript_v10.docx")
OUT = Path("/home/lysos/Desktop/Panagiotis/papers/RecChain/RecChain_manuscript_v11.docx")

doc = Document(str(SRC))

def find_heading(fragment):
    for p in doc.paragraphs:
        s = p.style.name if p.style else ''
        if 'Heading' in s and fragment in p.text:
            return p
    return None

def remove_heading(fragment):
    p = find_heading(fragment)
    if p:
        p._element.getparent().remove(p._element)
        return True
    return False

def rename_heading(fragment, new_text):
    p = find_heading(fragment)
    if p:
        p.clear()
        run = p.add_run(new_text)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return True
    return False

# ── 4+5: merge Oracle into Smart Contract chapter ────────────────────────────
rename_heading('4. Smart Contract',
               '4. Smart Contract Design and Oracle Integration')
remove_heading('5. Oracle-Based')
print("4+5 merged → 4. Smart Contract Design and Oracle Integration ✓")

# ── 6+7: merge Evaluation into IADM chapter ──────────────────────────────────
rename_heading('6. Intelligent Anomaly',
               '5. Anomaly Detection Module and Evaluation')
rename_heading('6.1 Detection',
               '5.1 Detection Methodology and Decision Logic')

remove_heading('7. Evaluation')
rename_heading('7.1 Anomaly Detection Evaluation',
               '5.2 Anomaly Detection Evaluation')
rename_heading('7.2 Sensitivity',
               '5.3 Sensitivity Analysis')
rename_heading('7.3 Gas Cost',
               '5.4 Gas Cost Analysis')
print("6+7 merged → 5. Anomaly Detection Module and Evaluation ✓")
print("  5.1–5.4 subsections renumbered ✓")

# ── 8+9+10: merge into single Implementation & Security chapter ───────────────
rename_heading('8. Implementation',
               '6. Implementation, Security Analysis, and Limitations')
remove_heading('9. Security')
remove_heading('10. Limitations')
rename_heading('10.1 Evaluation Scope',
               '6.1 Evaluation Scope and Detection Methodology Limitations')
rename_heading('10.2 Deployment Constraints',
               '6.2 Deployment Constraints and Future Research Directions')
print("8+9+10 merged → 6. Implementation, Security Analysis, and Limitations ✓")

# ── 11 → 7 ───────────────────────────────────────────────────────────────────
rename_heading('11. Conclusion', '7. Conclusion')
print("11. Conclusion → 7. Conclusion ✓")

# ── Re-apply black colour to all headings (clear() may lose colour) ───────────
for p in doc.paragraphs:
    s = p.style.name if p.style else ''
    if 'Heading' in s:
        for run in p.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save(str(OUT))
import os
print(f"\n✓ Saved: {OUT.name}  ({os.path.getsize(OUT)//1024} KB)")

# ── Print final structure ─────────────────────────────────────────────────────
doc2 = Document(str(OUT))
print("\nFinal section structure:")
for p in doc2.paragraphs:
    s = p.style.name if p.style else ''
    if 'Heading' in s and p.text.strip():
        indent = '    ' if '2' in s else ''
        print(f"  {indent}{p.text}")
