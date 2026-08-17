"""
fix_references.py
=================
Updates references section in v8 manuscript:
  1. Fixes [Khaqqi2018] DOI  (confirmed: 10.1016/j.apenergy.2017.10.070)
  2. Fixes [Worley2019] — corrects to 2018 IEEE Cybermatics conference
  3. Fixes [AIB23] — adds verified URL (Landsvirkjun / AIB announcement)
  4. Flags [Fahim2023] — marks as NEEDS MANUAL DOI (not found in search)
  5. Fixes [Nour2022] DOI   (correct: 10.1109/ACCESS.2022.3171227)
  6. Adds [Powerledger2018] reference
  7. Adds [WePower2018] reference
  8. Removes "→ ADD: Powerledger..." editor note
  9. Removes "→ Target: 18-22..." editor note (if still present)

Output: overwrites existing v8 manuscript
"""

from pathlib import Path
from docx import Document

SRC_OUT = Path("/home/lysos/Desktop/Panagiotis/papers/RecChain/RecChain_manuscript_v8_complete.docx")

doc = Document(str(SRC_OUT))

# ── Reference text replacements ───────────────────────────────────────────────

REPLACEMENTS = {
    # Khaqqi — fix DOI
    "[Khaqqi2018] Khaqqi K.N. et al. Blockchain-enabled emission trading. Applied Energy 209, 2018. [VERIFY DOI]":
        "[Khaqqi2018] Khaqqi K.N., Sikorski J.J., Hadinoto K., Kraft M. Incorporating seller/buyer reputation-based system in blockchain-enabled emission trading application. Applied Energy 209:8–19, 2018. DOI: 10.1016/j.apenergy.2017.10.070",

    # Worley — fix year and conference
    "[Worley2019] Worley C., Skjellum A. Blockchain tradeoffs for energy applications. IEEE IEMCON 2019. [VERIFY DOI]":
        "[Worley2018] Worley C., Skjellum A. Blockchain tradeoffs and challenges for current and emerging applications: generalization, fragmentation, sidechains, and scalability. In Proc. IEEE Cybermatics Congress (iThings/GreenCom/CPSCom/SmartData), Halifax, Canada, 2018. DOI: 10.1109/Cybermatics_2018.2018.00265",

    # AIB23 — add confirmed URL
    "[AIB23] AIB. Statement on Iceland GoO suspension. April 2023. [VERIFY: retrieve from aib-net.org]":
        "[AIB23] Association of Issuing Bodies (AIB). Suspension of Icelandic Guarantees of Origin exports following double-counting complaint. April 27, 2023. [Online]. Available: https://www.aib-net.org (see also: https://www.landsvirkjun.com/news/aib-suspends-all-exports-of-icelandic-gos)",

    # Fahim — flag clearly
    "[Fahim2023] Fahim M. et al. Anomaly Detection in Smart Grid Systems. IEEE, 2023. [VERIFY DOI]":
        "[Fahim2023] Fahim M. et al. Anomaly Detection in Smart Grid Systems. IEEE, 2023. [NOTE: DOI not confirmed — manual verification required via IEEE Xplore search for author 'Fahim' + 'anomaly detection' + 'smart grid' 2023]",

    # Nour — fix DOI
    "[Nour2022] Nour M. et al. Blockchain Applications in the Electricity Sector. IEEE Access 10, 2022. DOI: 10.1109/ACCESS.2022.3169069 [VERIFY]":
        "[Nour2022] Nour M., Chaves-Ávila J.P., Sánchez-Miralles Á. Review of blockchain potential applications in the electricity sector and challenges for large scale adoption. IEEE Access 10:47384–47418, 2022. DOI: 10.1109/ACCESS.2022.3171227",

    # Editor notes to remove
    "→ Target: 18–22. [VERIFY] = needs DOI confirmation before submission.":
        "",

    "→ ADD: Powerledger whitepaper, WePower docs, 1–2 IEEE Trans Smart Grid synthetic benchmark papers.":
        "",
}

changed = []
for para in doc.paragraphs:
    text = para.text
    for old, new in REPLACEMENTS.items():
        if old in text:
            if new == "":
                # Remove the paragraph
                el = para._element
                el.getparent().remove(el)
                changed.append(f"REMOVED: {old[:60]}...")
            else:
                # Replace text
                para.clear()
                para.add_run(new)
                changed.append(f"UPDATED: {old[:50]}...")
            break

# ── Add missing references before the closing of the References section ───────
# Find the last reference paragraph to insert after
last_ref = None
for para in doc.paragraphs:
    if para.text.strip().startswith('[') and para.text.strip()[1:4].isupper() or \
       (para.text.strip().startswith('[') and any(c.isupper() for c in para.text[1:5])):
        last_ref = para

# Add Powerledger and WePower after the last reference
if last_ref:
    powerledger = "[Powerledger2018] Powerledger Pty Ltd. Powerledger Whitepaper: A Peer-to-Peer Renewable Energy Trading Platform. 2017/2018. [Online]. Available: https://powerledger.io/company/power-ledger-whitepaper/"
    wepower = "[WePower2018] WePower. WePower Whitepaper: A Global Green Energy Network Powered by Smart Contracts. 2018. [Online]. Available: https://medium.com/wepower/updated-wepower-whitepaper-is-now-available-c2c9468918c4"

    for new_ref in [powerledger, wepower]:
        new_para = doc.add_paragraph()
        new_para.add_run(new_ref)
        last_ref._element.addnext(new_para._element)
        changed.append(f"ADDED: {new_ref[:60]}...")

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save(str(SRC_OUT))

print("Reference updates applied:")
for c in changed:
    print(f"  {c}")

# ── Count final references ────────────────────────────────────────────────────
doc2 = Document(str(SRC_OUT))
refs = [p.text for p in doc2.paragraphs if p.text.strip().startswith('[') and ']' in p.text]
print(f"\nTotal references: {len(refs)}")
for r in refs:
    print(f"  {r[:90]}")
