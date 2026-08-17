"""
update_figures_v10.py
=====================
Takes RecChain_manuscript_v9_formatted.docx (with manual edits) as source.
Replaces ONLY the 3 redrawn figures (Fig 1, Fig 2, Fig 7) with the new
clean matplotlib versions.
All text, formatting, and manual edits are preserved.
Output: RecChain_manuscript_v10.docx
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn

SRC = Path("/home/lysos/Desktop/Panagiotis/papers/RecChain/RecChain_manuscript_v9_formatted.docx")
OUT = Path("/home/lysos/Desktop/Panagiotis/papers/RecChain/RecChain_manuscript_v10.docx")
FIGS = Path("/home/lysos/Desktop/Panagiotis/Projects/REC_Dapp_Ethereum/evaluation/results/figures")

doc = Document(str(SRC))
paras = doc.paragraphs

def get_caption(para_idx):
    try: return paras[para_idx + 1].text
    except: return ''

# Identify which paragraph index holds each figure by reading caption after it
fig_map = {}
for i, p in enumerate(paras):
    blips = p._element.findall('.//' + qn('a:blip'),
        namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
    if blips:
        cap = get_caption(i)
        if 'Figure 1.' in cap:  fig_map[1] = i
        elif 'Figure 2.' in cap: fig_map[2] = i
        elif 'Figure 7.' in cap: fig_map[7] = i

print("Figures to replace:")
for k, v in sorted(fig_map.items()):
    print(f"  Fig {k} → paragraph {v}  |  {get_caption(v)[:60]}")

def replace_image_in_para(para, png_path, width_in=5.5):
    """Remove old drawing and insert new image in same paragraph."""
    # Remove all existing drawing/picture XML elements
    for child in list(para._element):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'r':
            # Check if this run has a drawing
            drawings = child.findall('.//' + qn('w:drawing'))
            if drawings:
                para._element.remove(child)
    # Also remove bare <w:drawing> elements
    for drawing in para._element.findall('.//' + qn('w:drawing')):
        parent = drawing.getparent()
        parent.getparent().remove(parent)
    # Add new picture
    run = para.add_run()
    run.add_picture(str(png_path), width=Inches(width_in))

# Replace Fig 1
if 1 in fig_map:
    replace_image_in_para(paras[fig_map[1]], FIGS / 'fig1_architecture.png', 5.5)
    print("Fig 1 replaced ✓")

# Replace Fig 2
if 2 in fig_map:
    replace_image_in_para(paras[fig_map[2]], FIGS / 'fig2_oracle_sequence.png', 5.5)
    print("Fig 2 replaced ✓")

# Replace Fig 7
if 7 in fig_map:
    replace_image_in_para(paras[fig_map[7]], FIGS / 'fig7_threat_model.png', 5.5)
    print("Fig 7 replaced ✓")

doc.save(str(OUT))
import os
size = os.path.getsize(OUT) // 1024
print(f"\n✓ Saved: {OUT.name}  ({size} KB)")
print(f"  All manual edits from v9 preserved.")
print(f"  Only Fig 1, Fig 2, Fig 7 images were replaced.")

# Verify image count
doc2 = Document(str(OUT))
print(f"  Total inline images: {len(doc2.inline_shapes)}")
